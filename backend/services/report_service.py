import io
from datetime import datetime, date
from typing import List, Dict, Any

# openpyxl for Excel reports
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab for PDF reports
from reportlab.lib.pagesizes import letter, A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas to dynamically calculate the total number of pages and
    print 'Page X of Y' in the footer.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#4B5563"))
        
        # Draw a footer line
        self.setStrokeColor(colors.HexColor("#D1D5DB"))
        self.setLineWidth(0.5)
        # Handle landscape vs portrait dimensions
        width, height = self._pagesize
        self.line(36, 45, width - 36, 45)
        
        # Footer text
        footer_text = f"Jharkhand Bijli Vitran Nigam Limited | Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        page_text = f"Page {self._pageNumber} of {page_count}"
        
        self.drawString(36, 30, footer_text)
        self.drawRightString(width - 36, 30, page_text)
        self.restoreState()


def generate_excel_report(data: List[Dict[str, Any]], title: str, headers: List[str], keys: List[str]) -> bytes:
    """
    Generate an Excel sheet with formatted headers, auto-adjusted columns, and border styles.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True

    # Style definitions
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A5F")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Title block
    ws.merge_cells('A1:G1')
    ws['A1'] = title.upper()
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    
    ws['A2'] = f"Generated Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = Font(name="Calibri", size=10, italic=True)
    ws.row_dimensions[2].height = 18

    # Table headers
    start_row = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[start_row].height = 25

    # Data Rows
    current_row = start_row + 1
    for row_data in data:
        for col_idx, key in enumerate(keys, 1):
            val = row_data.get(key, "")
            # Convert date/datetime to string for presentation
            if isinstance(val, (datetime, date)):
                val = val.strftime('%Y-%m-%d')
            elif val is None:
                val = ""
                
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            
            # Format numbers
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
        current_row += 1

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < start_row:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Output to bytes
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def generate_pdf_report(
    data: List[Dict[str, Any]], 
    title: str, 
    headers: List[str], 
    keys: List[str], 
    filters_desc: str = ""
) -> bytes:
    """
    Generate a ReportLab PDF document using the Jharkhand electricity office letterhead layout.
    Automatically chooses Landscape orientation for reports with more than 6 columns.
    """
    buffer = io.BytesIO()
    
    # Decide page setup based on width
    is_landscape = len(headers) > 6
    page_size = landscape(A4) if is_landscape else A4
    
    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=36,
        rightMargin=36,
        topMargin=40,
        bottomMargin=60
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor("#1e3a5f"),
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        textColor=colors.HexColor("#b58900"),
        spaceAfter=15
    )
    
    filter_style = ParagraphStyle(
        'DocFilter',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        textColor=colors.HexColor("#4B5563"),
        spaceAfter=10
    )

    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10
    )
    
    cell_header_style = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1 # Centered
    )

    story = []

    # Letterhead Banner representation
    letterhead_text = "<b>JHARKHAND BIJLI VITRAN NIGAM LIMITED</b><br/>Government of Jharkhand Undertaking"
    lh_style = ParagraphStyle(
        'LetterHead',
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=14,
        textColor=colors.HexColor("#1e3a5f")
    )
    
    # Draw letterhead logo block placeholder style
    logo_data = [[
        Paragraph(letterhead_text, lh_style),
        Paragraph("<font size=8>Contact: contact@jbvnl.co.in<br/>Web: www.jbvnl.co.in</font>", ParagraphStyle('LHR', alignment=2))
    ]]
    lh_table = Table(logo_data, colWidths=[doc.width * 0.7, doc.width * 0.3])
    lh_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
        ('LINEBELOW', (0,0), (-1,-1), 1.5, colors.HexColor("#1e3a5f")),
    ]))
    
    story.append(lh_table)
    story.append(Spacer(1, 15))

    # Title & Filters
    story.append(Paragraph(title.upper(), title_style))
    story.append(Paragraph("PROJECT MONITORING & PROGRESS DOCUMENTATION", subtitle_style))
    
    if filters_desc:
        story.append(Paragraph(f"<b>Applied Filters:</b> {filters_desc}", filter_style))

    # Build Table
    table_data = []
    
    # Wrap headers in Paragraphs for autowrap
    header_row = [Paragraph(h, cell_header_style) for h in headers]
    table_data.append(header_row)
    
    # Build rows
    for item in data:
        row = []
        for key in keys:
            val = item.get(key, "")
            if isinstance(val, (datetime, date)):
                val_str = val.strftime('%Y-%m-%d')
            elif val is None:
                val_str = ""
            elif isinstance(val, (int, float)):
                val_str = f"{val:,.2f}"
            else:
                val_str = str(val)
                
            row.append(Paragraph(val_str, cell_style))
        table_data.append(row)

    # Set column widths proportionally
    col_count = len(headers)
    col_width = doc.width / col_count
    
    t = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    
    # Styling Table
    t_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
    ])
    t.setStyle(t_style)
    story.append(t)

    # Build document
    doc.build(story, canvasmaker=NumberedCanvas)
    
    buffer.seek(0)
    return buffer.getvalue()
